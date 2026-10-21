"""Spreadsheet exports for the admin grids.

Replaces the reference stack's ExcelJS/csv-writer usage. Returns bytes so the
controller can stream them with the right content-disposition header.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", start_color="1F5F3F", end_color="1F5F3F")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _stringify(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.strftime(
            "%Y-%m-%d %H:%M" if isinstance(value, datetime) else "%Y-%m-%d"
        )
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (dict, list)):
        import json

        return json.dumps(value, default=str)
    return "" if value is None else value


def to_xlsx(
    rows: Sequence[dict[str, Any]],
    columns: Sequence[tuple[str, str]],
    sheet_name: str = "Export",
) -> bytes:
    """``columns`` is ``[(key, header_label), ...]`` in display order."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31] or "Export"

    for index, (_, label) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, row in enumerate(rows, start=2):
        for col_index, (key, _) in enumerate(columns, start=1):
            sheet.cell(row=row_index, column=col_index, value=_stringify(row.get(key)))

    # Size columns to their content, within sane bounds.
    for col_index, (key, label) in enumerate(columns, start=1):
        widest = max(
            [len(str(label))]
            + [len(str(_stringify(row.get(key)))) for row in rows[:200]]
            or [10]
        )
        sheet.column_dimensions[get_column_letter(col_index)].width = min(
            max(widest + 2, 12), 55
        )

    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_csv(rows: Sequence[dict[str, Any]], columns: Sequence[tuple[str, str]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow([label for _, label in columns])
    for row in rows:
        writer.writerow([_stringify(row.get(key)) for key, _ in columns])
    # BOM so Excel opens UTF-8 (Amharic names included) without mangling.
    return buffer.getvalue().encode("utf-8-sig")


def export(
    rows: Sequence[dict[str, Any]],
    columns: Sequence[tuple[str, str]],
    fmt: str = "xlsx",
    sheet_name: str = "Export",
) -> tuple[bytes, str, str]:
    """Returns ``(content, media_type, file_extension)``."""
    if fmt == "csv":
        return to_csv(rows, columns), "text/csv; charset=utf-8", "csv"
    return (
        to_xlsx(rows, columns, sheet_name),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    )


# Column sets reused by the export endpoints.
PARTICIPANT_COLUMNS: list[tuple[str, str]] = [
    ("participant_id", "ID"),
    ("display_name", "Name"),
    ("company_name", "Company"),
    ("profile_type", "Profile Type"),
    ("email_id", "Email"),
    ("mobile_number", "Mobile"),
    ("country_code", "Country"),
    ("admin_area_code_l1", "Region"),
    ("category_id", "Business Category"),
    ("participant_type_id", "Network Role"),
    ("is_email_verified", "Email Verified"),
    ("is_sandbox_approved", "Sandbox"),
    ("is_uat_approved", "UAT"),
    ("is_prod_approved", "Production"),
    ("is_rejected", "Rejected"),
    ("rejection_reason", "Rejection Reason"),
    ("register_date", "Registered On"),
    ("approved_date", "Approved On"),
]

API_LOG_COLUMNS: list[tuple[str, str]] = [
    ("log_timestamp", "Timestamp"),
    ("transaction_id", "Transaction"),
    ("participant_id", "User"),
    ("environment", "Environment"),
    ("method", "Method"),
    ("endpoint_url", "Endpoint"),
    ("status_code", "Status"),
    ("latency_ms", "Latency (ms)"),
    ("client_ip", "Client IP"),
    ("error_message", "Error"),
]

USER_ACTION_COLUMNS: list[tuple[str, str]] = [
    ("timestamp", "Timestamp"),
    ("actor_type", "Actor Type"),
    ("actor_name", "Actor"),
    ("action", "Action"),
    ("entity", "Entity"),
    ("entity_id", "Entity ID"),
    ("description", "Description"),
    ("ip_address", "IP Address"),
]
